"""Tests for extracting per-destination volumes from a Royal Mail carrier sheet."""

import os
import sys
import shutil
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook

from carriers.royalmail import RoyalMailCarrier


def build_sheet(path, rows, po_number=123456, carrier_name="Royal Mail International 2026"):
    """
    Write a carrier sheet in the real layout: name in B3, PO in B4,
    headers on row 8, data from row 9.
    """
    wb = Workbook()
    ws = wb.active
    ws['B3'] = carrier_name
    ws['B4'] = po_number
    for col, header in enumerate(['Country', 'Service', 'Postcode', 'Format', 'Items', 'Weight (KG)'], start=1):
        ws.cell(row=8, column=col, value=header)
    for offset, (country, fmt, items, weight) in enumerate(rows):
        row = 9 + offset
        ws.cell(row=row, column=1, value=country)
        ws.cell(row=row, column=4, value=fmt)
        ws.cell(row=row, column=5, value=items)
        ws.cell(row=row, column=6, value=weight)
    wb.save(path)
    wb.close()


class RoyalMailSheetTests(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.sheet = os.path.join(self.tmp, 'carrier_sheet.xlsx')
        self.carrier = RoyalMailCarrier()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def process(self, rows, **kwargs):
        build_sheet(self.sheet, rows, **kwargs)
        return self.carrier.extract_data(self.sheet)

    def lines_of(self, data):
        return sorted(
            (l.country, l.format_type, l.items, round(l.weight_kg, 3))
            for l in data.lines
        )

    def test_each_destination_and_format_becomes_its_own_line(self):
        data = self.process([
            ('Ireland', 'Letters', 100, 5.0),
            ('Netherlands', 'Letters', 200, 8.0),
            ('Canada', 'Flats', 50, 12.5),
        ])
        self.assertEqual(self.lines_of(data), [
            ('Canada', 'Flats', 50, 12.5),
            ('Ireland', 'Letters', 100, 5.0),
            ('Netherlands', 'Letters', 200, 8.0),
        ])

    def test_repeated_rows_for_one_destination_are_summed(self):
        data = self.process([
            ('France', 'Letters', 100, 4.0),
            ('France', 'Letters', 50, 2.0),
        ])
        self.assertEqual(self.lines_of(data), [('France', 'Letters', 150, 6.0)])

    def test_country_aliases_are_resolved_to_one_line(self):
        data = self.process([
            ('Republic of Ireland', 'Letters', 100, 4.0),
            ('ROI', 'Letters', 25, 1.0),
        ])
        self.assertEqual(self.lines_of(data), [('Ireland', 'Letters', 125, 5.0)])

    def test_letters_and_flats_for_one_country_stay_separate(self):
        data = self.process([
            ('Germany', 'Letters', 100, 4.0),
            ('Germany', 'Flats', 10, 6.0),
        ])
        self.assertEqual(self.lines_of(data), [
            ('Germany', 'Flats', 10, 6.0),
            ('Germany', 'Letters', 100, 4.0),
        ])

    def test_an_unsupported_destination_fails_loudly(self):
        with self.assertRaises(ValueError) as ctx:
            self.process([
                ('Ireland', 'Letters', 100, 4.0),
                ('Spain', 'Letters', 60, 3.0),
            ])
        message = str(ctx.exception)
        self.assertIn('Spain', message)

    def test_an_unrecognised_format_fails_loudly(self):
        with self.assertRaises(ValueError) as ctx:
            self.process([('Ireland', 'Parcels', 100, 40.0)])
        self.assertIn('Parcels', str(ctx.exception))

    def test_blank_rows_are_skipped(self):
        data = self.process([
            ('Ireland', 'Letters', 100, 4.0),
            (None, None, None, None),
            ('Italy', 'Flats', 10, 3.0),
        ])
        self.assertEqual(len(data.lines), 2)

    def test_a_float_po_number_is_read_as_a_whole_number(self):
        data = self.process([('Ireland', 'Letters', 1, 0.1)], po_number=987654.0)
        self.assertEqual(data.po_number, '987654')

    def test_aggregate_totals_stay_available_for_the_processing_log(self):
        data = self.process([
            ('Ireland', 'Letters', 100, 5.0),
            ('Italy', 'Letters', 50, 2.0),
            ('Canada', 'Flats', 20, 8.0),
        ])
        self.assertEqual(data.letters_items, 150)
        self.assertAlmostEqual(data.letters_weight, 7.0)
        self.assertEqual(data.flats_items, 20)
        self.assertAlmostEqual(data.flats_weight, 8.0)

    def test_extracting_data_writes_no_files(self):
        build_sheet(self.sheet, [('Ireland', 'Letters', 1, 0.1)])
        before = set(os.listdir(self.tmp))
        self.carrier.extract_data(self.sheet)
        self.assertEqual(set(os.listdir(self.tmp)), before)


if __name__ == '__main__':
    unittest.main()
